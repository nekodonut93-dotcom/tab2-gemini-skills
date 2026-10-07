from __future__ import annotations

import json
from pathlib import Path
import zipfile
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "docs" / "generated" / "TAB2_GEMINI_SKILLS_KIEN_TRUC_VAN_HANH_VI.docx"

def read(relative: str) -> str:
    return (ROOT / relative).read_text(encoding="utf-8")

def add_source_note(doc, relative: str) -> None:
    doc.add_paragraph(f"Nguồn trong repository: {relative}", style="Caption")

def add_skill(doc, relative: str, label: str) -> None:
    path = ROOT / relative
    doc.add_heading(label, level=2)
    text = path.read_text(encoding="utf-8")
    for line in text.splitlines():
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=3)
        elif line.strip().startswith("- "):
            doc.add_paragraph(line.strip()[2:], style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line.strip())
    add_source_note(doc, relative)

def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("Kiến trúc và vận hành Tab 2 Gemini Skills", 0)
    doc.add_paragraph("Tài liệu này được tổng hợp chỉ từ repository tab2-gemini-skills. Repository là gói tham chiếu portable; production vẫn là source of truth và không giao runtime authority cho Gemini.")
    skills = [
        ("00 orchestrator", "skills/00-orchestrator/SKILL.md"),
        ("01 story-memory", "skills/01-story-memory/SKILL.md"),
        ("02 continuity", "skills/02-continuity/SKILL.md"),
        ("03 writer", "skills/03-writer/SKILL.md"),
        ("04 enrichment", "skills/04-enrichment/SKILL.md"),
        ("04b visual-completion", "skills/04b-visual-completion/SKILL.md"),
        ("04c action-interpreter", "skills/04c-action-interpreter/SKILL.md"),
        ("04d action-choreography", "skills/04d-action-choreography/SKILL.md"),
        ("04e chibi-performance", "skills/04e-chibi-performance/SKILL.md"),
        ("05 director", "skills/05-director/SKILL.md"),
        ("06 prompt-engine", "skills/06-prompt-engine/SKILL.md"),
        ("07 quality-audit", "skills/07-quality-audit/SKILL.md"),
        ("08 seedance", "skills/08-seedance/SKILL.md"),
        ("09 style-resolution", "skills/09-style-resolution/SKILL.md"),
    ]
    topics = [
        ("1. Tổng quan hệ thống", "README.md"),
        ("2. Sơ đồ pipeline", "00_CORE/WORKFLOW.md"),
        ("3. Cấu trúc repository", "README.md"),
        ("4. Giải thích từng folder", "README.md"),
        ("5. Giải thích từng skill", "__SKILLS__"),
        ("6. Authority", "00_CORE/DATA_AUTHORITY.md"),
        ("7. Quy trình đưa tiểu thuyết vào", "skills/00-orchestrator/SKILL.md"),
        ("8. Style questionnaire", "skills/09-style-resolution/style_questionnaire.yaml"),
        ("9. Story Memory", "skills/01-story-memory/SKILL.md"),
        ("10. Continuity", "skills/02-continuity/SKILL.md"),
        ("11. Writer WHAT", "skills/03-writer/SKILL.md"),
        ("12. Visual Completion", "skills/04b-visual-completion/SKILL.md"),
        ("13. Safe vs forbidden inference", "00_CORE/CREATIVE_BOUNDARY.md"),
        ("14. Action Interpreter", "skills/04c-action-interpreter/SKILL.md"),
        ("15. Action Choreography", "skills/04d-action-choreography/SKILL.md"),
        ("16. Chibi performance", "skills/04e-chibi-performance/SKILL.md"),
        ("17. Language policy", "language/LANGUAGE_POLICY.md"),
        ("18. Director HOW", "skills/05-director/SKILL.md"),
        ("19. Seedance", "skills/08-seedance/SKILL.md"),
        ("20. Dola", "seedance/platforms/dola.yaml"),
        ("21. Character references", "seedance/reference_binding/multi_reference.yaml"),
        ("22. @Image handling", "seedance/platforms/dola.yaml"),
        ("23. Timeline 5/10/15s", "seedance/timing/README.md"),
        ("24. Chibi/comedy", "seedance/vfx/comedic.yaml"),
        ("25. Anime", "skills/09-style-resolution/style_presets.yaml"),
        ("26. Chinese 3D", "skills/09-style-resolution/style_presets.yaml"),
        ("27. Wuxia/Xianxia", "seedance/vfx/wuxia.yaml"),
        ("28. VFX", "seedance/vfx/fantasy.yaml"),
        ("29. Prompt structure", "seedance/PROMPT_CONTRACT.yaml"),
        ("30. Quality Audit", "skills/07-quality-audit/SKILL.md"),
        ("31. Sheet export", "seedance/sheet_export/README.md"),
        ("32. Ví dụ end-to-end", "benchmarks/END_TO_END_PORTABLE_EXAMPLE.yaml"),
        ("33. Known unknowns", "audit/production_capability_audit.json"),
        ("34. Missing capabilities", "audit/production_capability_audit.json"),
        ("35. Checklist bổ sung", "audit/USER_GAP_CHECKLIST.md"),
        ("36. Ngôn ngữ prompt và ngôn ngữ thoại", "language/LANGUAGE_POLICY.md"),
        ("37. Prompt tiếng Anh / thoại tiếng Việt", "language/examples/vi_source_en_prompt_vi_speech.json"),
        ("38. Thoại nội tâm", "language/inner_monologue_policy.yaml"),
        ("39. Chibi nội tâm", "skills/04e-chibi-performance/inner_monologue/routing.yaml"),
        ("40. Chibi nói chuyện", "skills/04e-chibi-performance/inner_monologue/spoken_chibi.yaml"),
        ("41. Chibi cử động tượng trưng", "skills/04e-chibi-performance/inner_monologue/symbolic_action.yaml"),
        ("42. Đấm/đá thật và đấm/đá chibi", "skills/04e-chibi-performance/grammars/fight_pantomime.yaml"),
        ("43. Khóc thật và khóc chibi", "skills/04e-chibi-performance/grammars/crying.yaml"),
        ("44. Ranh giới thế giới thật / non-diegetic", "skills/04e-chibi-performance/constraints/diegetic_boundary.yaml"),
        ("45. Ví dụ hoàn chỉnh", "benchmarks/vi_chibi_inner_monologue_example.yaml"),
    ]
    doc.add_heading("Nội dung", level=1)
    for heading, relative in topics:
        doc.add_heading(heading, level=1)
        if relative == "__SKILLS__":
            for label, skill_path in skills:
                add_skill(doc, skill_path, label)
            continue
        if relative.endswith(".json"):
            data = json.loads(read(relative))
            if heading.startswith("33."):
                for item in data["capabilities"]:
                    if item["status"] != "VERIFIED":
                        doc.add_paragraph(f"{item['id']} — {item['status']}: {'; '.join(item['known_gaps'])}", style="List Bullet")
            elif heading.startswith("34."):
                for item in data["capabilities"]:
                    if item["status"] in {"NOT_IMPLEMENTED", "MISSING", "UNKNOWN"}:
                        doc.add_paragraph(f"{item['id']}: {item['what_is_evidenced']} Missing: {'; '.join(item['known_gaps'])}", style="List Bullet")
            else:
                doc.add_paragraph(json.dumps(data, ensure_ascii=False, indent=2))
        else:
            text = read(relative)
            for line in text.splitlines():
                if line.startswith("#"):
                    continue
                if line.strip().startswith("- "):
                    doc.add_paragraph(line.strip()[2:], style="List Bullet")
                elif line.strip():
                    doc.add_paragraph(line.strip())
        add_source_note(doc, relative)
    audit = json.loads(read("audit/production_capability_audit.json"))
    doc.add_heading("Bảng trạng thái thành phần", level=1)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Shading Accent 1"
    headers = ["COMPONENT", "STATUS", "SOURCE", "WHAT IT DOES", "MISSING", "RECOMMENDED NEXT ACTION"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for item in audit["capabilities"]:
        row = table.add_row().cells
        values = [
            item["id"], item["status"], "; ".join(item["production_sources"]),
            item["what_is_evidenced"], "; ".join(item["known_gaps"]),
            "Kiểm tra bằng chứng và gap; chỉ bổ sung khi có production evidence hoặc quyết định migration rõ ràng.",
        ]
        for cell, value in zip(row, values):
            cell.text = value
    doc.core_properties.title = "Kiến trúc vận hành Tab 2 Gemini Skills"
    doc.core_properties.author = "Tab 2 Gemini Skills repository"
    doc.core_properties.created = __import__("datetime").datetime(2026, 10, 7)
    doc.core_properties.modified = __import__("datetime").datetime(2026, 10, 7)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    normalize_docx_zip(OUT)
    manifest_path = ROOT / "SOURCE_MANIFEST.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    source_production_files = sorted({
        source
        for metadata in manifest["files"].values()
        for source in metadata["source_production_files"]
    })
    manifest["files"]["docs/generated/TAB2_GEMINI_SKILLS_KIEN_TRUC_VAN_HANH_VI.docx"] = {
        "source_production_files": source_production_files,
        "source_skill_version": None,
        "schema_version": "VIETNAMESE_TAB2_GEMINI_GUIDE_V1",
        "distribution_version": manifest["distribution_version"],
        "status": "GENERATED_FROM_VALIDATED_PORTABLE_REPOSITORY",
        "known_gaps": ["Documents describe portable package behavior only; production-only unknowns are explicitly listed in the document."],
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(OUT)

def normalize_docx_zip(path: Path) -> None:
    fixed_time = (2020, 1, 1, 0, 0, 0)
    temp_path = path.with_suffix(".normalized.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp_path, "w") as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, date_time=fixed_time)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = source.getinfo(name).external_attr
            target.writestr(info, source.read(name))
    temp_path.replace(path)

if __name__ == "__main__":
    main()
